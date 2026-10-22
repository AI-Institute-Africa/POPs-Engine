"""Every page of PDFs, Word documents and HEIC photos."""

import io
import uuid
import zipfile

import pymupdf
import pytest
from docx import Document
from PIL import Image
from sqlalchemy import select

from app.intake.files import RejectedFile, prepare
from app.models import VerificationRun
from tests.conftest import ecocash_fields, login, make_receipt
from tests.test_flow import TXN, _import, _obligation, _stmt_row

RECEIPT_TEXT = [
    "EcoCash",
    "Payment Successful",
    f"Txn ID: {TXN}",
    "Amount: USD 40.00",
    "Paid to: COGNIFY MAIN",
    "21/10/2026 14:32",
]


def _pdf(pages: list[list[str]]) -> bytes:
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page()
        for n, line in enumerate(lines):
            page.insert_text((72, 72 + 24 * n), line, fontsize=14)
    return doc.tobytes()


def _docx(lines: list[str], image: bytes | None = None) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if image:
        doc.add_picture(io.BytesIO(image))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(client, h, ob, data, name, ctype):
    r = client.post(
        "/claims",
        headers=h,
        files={"file": (name, data, ctype)},
        data={
            "obligation_id": ob["obligation_id"],
            "claimant_ref": "learner-1",
            "channel": "ecocash",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


def _internal(client, claim_id):
    return client.get(f"/claims/{claim_id}", headers=login(client, "reviewer")).json()


def test_every_pdf_page_is_read():
    p = prepare(
        _pdf([["Cover letter from the payer"], RECEIPT_TEXT, ["Terms and conditions"]]), 10**8
    )
    assert len(p.pages) == 3 and TXN in p.page_texts[1]
    assert [round(y, 3) for y, _ in p.page_layout] == [0.0, 0.333, 0.667]
    assert p.meta["pdf"]["pages_read"] == 3


def test_pdf_receipt_on_page_two_is_matched(client, session):
    sub, rev = login(client, "submitter"), login(client, "reviewer")
    ob = _obligation(client, sub)
    _import(client, rev, [_stmt_row()])
    claim = _upload(
        client,
        sub,
        ob,
        _pdf([["Proof of payment for my daughter"], RECEIPT_TEXT]),
        "pop.pdf",
        "application/pdf",
    )
    assert claim["state"] == "FINANCIALLY_CONFIRMED", claim["reason_codes"]
    run = session.scalars(
        select(VerificationRun).where(VerificationRun.claim_id == uuid.UUID(claim["claim_id"]))
    ).one()
    assert "text_layer" in run.model_versions
    extraction = next(c for c in run.checks if c["check"] == "extraction")
    assert extraction["pages"] == 2


def test_scanned_pdf_without_text_is_read_from_its_pixels(client):
    sub = login(client, "submitter")
    ob = _obligation(client, sub)
    img = Image.open(io.BytesIO(make_receipt({"amount": "40.00"})))
    buf = io.BytesIO()
    img.save(buf, format="PDF")
    claim = _upload(client, sub, ob, buf.getvalue(), "scan.pdf", "application/pdf")
    assert claim["state"] in ("PENDING_REVIEW", "NEEDS_INFORMATION")  # no text layer to read


def test_word_document_is_read_and_flagged_as_editable(client):
    sub, rev = login(client, "submitter"), login(client, "reviewer")
    ob = _obligation(client, sub)
    _import(client, rev, [_stmt_row()])
    claim = _upload(
        client,
        sub,
        ob,
        _docx(RECEIPT_TEXT),
        "pop.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    c = _internal(client, claim["claim_id"])
    assert c["state"] == "FINANCIALLY_CONFIRMED"
    assert "EDITABLE_DOCUMENT" in {s["code"] for s in c["risk_signals"]}


def test_word_document_with_embedded_screenshot_reads_every_image():
    shot = make_receipt({"amount": "40.00"}, size=(600, 900))
    p = prepare(_docx(["Please find my proof of payment attached."], image=shot), 10**8)
    assert len(p.pages) == 2  # rendered text + the screenshot
    assert p.meta["embedded_images"] == 1


def test_macro_documents_are_refused():
    data = _docx(RECEIPT_TEXT)
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/vbaProject.bin", b"\x00" * 64)
    with pytest.raises(RejectedFile) as e:
        prepare(buf.getvalue(), 10**8)
    assert e.value.code == "MACRO_DOCUMENT"


def test_other_zip_files_are_not_documents():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "hello")
    with pytest.raises(RejectedFile) as e:
        prepare(buf.getvalue(), 10**8)
    assert e.value.code == "UNSUPPORTED_TYPE"


def test_heic_photos_are_accepted():
    pillow_heif = pytest.importorskip("pillow_heif")
    img = Image.new("RGB", (640, 960), "white")
    buf = io.BytesIO()
    try:
        pillow_heif.from_pillow(img).save(buf, format="HEIF")
    except Exception as exc:  # noqa: BLE001 - encoder may be missing on some builds
        pytest.skip(f"HEIF encoder unavailable: {exc}")
    p = prepare(buf.getvalue(), 10**8)
    assert p.content_type == "image/heic" and p.image.size == (640, 960)


def _pdf_statement(rows: list[list[str]]) -> bytes:
    """A digital bank statement: a ruled table, split over two pages."""
    doc = pymupdf.open()
    xs = [40, 140, 260, 400, 460, 510, 560]
    half = 2  # header + one line, the rest continue on page two without a header
    for chunk in (rows[:half], rows[half:]):
        page = doc.new_page()
        page.insert_text((40, 50), "Statement of account MAIN", fontsize=12)
        for r, row in enumerate(chunk):
            for c, cell in enumerate(row):
                page.insert_text((xs[c] + 2, 94 + r * 20), cell, fontsize=7)
        for r in range(len(chunk) + 1):
            page.draw_line((xs[0], 80 + r * 20), (xs[-1], 80 + r * 20))
        for x in xs:
            page.draw_line((x, 80), (x, 80 + len(chunk) * 20))
    return doc.tobytes()


def test_pdf_bank_statement_is_imported_and_matches(client):
    sub, rev = login(client, "submitter"), login(client, "reviewer")
    ob = _obligation(client, sub)
    claim = _upload(client, sub, ob, make_receipt(ecocash_fields()), "pop.png", "image/png")
    assert claim["state"] == "AWAITING_FINANCIAL_CONFIRMATION", claim["reason_codes"]
    pdf = _pdf_statement(
        [
            ["DATE", "Narrative", "Reference", "Credit", "Debit", "Currency"],
            ["20/10/2026 08:00", "Bank charges", "", "", "2.00", "USD"],
            ["21/10/2026 14:35", "EcoCash payment", TXN, "40.00", "", "USD"],
            ["21/10/2026 16:00", "Transfer in", "FT26294XYZ", "15.00", "", "USD"],
        ]
    )
    r = client.post(
        "/credits/imports",
        headers=rev,
        data={"account_alias": "MAIN", "profile": "generic"},
        files={"file": ("stmt.pdf", pdf, "application/pdf")},
    )
    assert r.status_code == 201, r.text
    assert r.json()["created"] == 2 and r.json()["rows"] == 3, r.json()
    assert _internal(client, claim["claim_id"])["state"] == "FINANCIALLY_CONFIRMED"


def test_pdf_statement_without_a_table_is_refused(client):
    rev = login(client, "reviewer")
    r = client.post(
        "/credits/imports",
        headers=rev,
        data={"account_alias": "MAIN", "profile": "generic"},
        files={"file": ("stmt.pdf", _pdf([["Dear customer, thank you"]]), "application/pdf")},
    )
    assert r.status_code == 422 and r.json()["detail"]["code"] == "NO_STATEMENT_TABLE"
